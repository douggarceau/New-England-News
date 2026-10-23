import json
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'output'
PACKAGE = OUT / 'Rhode-Island-Hurricane-November-2026'
STORIES = PACKAGE / 'Stories'
OUT.mkdir(parents=True, exist_ok=True)
STORIES.mkdir(parents=True, exist_ok=True)


def normalize(data, filename=''):
    result = dict(data)
    result.setdefault('slug', re.sub(r'[^a-z0-9]+', '-', data['title'].lower()).strip('-'))
    section = 'Culture'
    if 'national' in filename:
        section = 'National and Technology'
    elif 'celebrities' in filename:
        section = 'People'
    elif 'food' in filename:
        section = 'Food'
    result.setdefault('section', section)
    result.setdefault('dek', '')
    result['body'] = result.get('body', result.get('paragraphs', []))
    result['body'] = [({'text': p, 'sources': []} if isinstance(p, str) else dict(p)) for p in result['body']]
    for p in result['body']:
        p.setdefault('sources', [])
    result['source_notes'] = result.get('source_notes', result.get('sources', []))
    for i, source in enumerate(result['source_notes'], 1):
        source.setdefault('id', i)
        source.setdefault('outlet', source.get('publisher', 'Source'))
        source.setdefault('title', source.get('name', source.get('url', 'Source')))
        source['date'] = source.get('date') or source.get('published') or 'Accessed October 23, 2026'
    result['word_count'] = len(' '.join(p['text'] for p in result['body']).split())
    return result


taylor = normalize(json.loads((ROOT.parent / 'ri-cover-story' / 'article.json').read_text()))
taylor['slug'] = 'taylor-swift-ashley-taunton'
taylor['section'] = 'Cover Story'
taylor['reporting_limits'] = [
    'The rescue occurred in July and the donation was reported in August. The November issue must not present either as a new event.',
    'The most recent recovery interview cited is dated September 1. The fundraising total is a dated October 23 snapshot.',
    'No new interview with the family or Taylor Swift was conducted. Irina Vladi is the cover model and is not Ashley Taunton.'
]
articles = [taylor]
for path in sorted((ROOT / 'research').rglob('*.json')):
    data = json.loads(path.read_text())
    candidates = data if isinstance(data, list) else data.get('articles', [data])
    for candidate in candidates:
        if isinstance(candidate, dict) and candidate.get('title') and (candidate.get('body') or candidate.get('paragraphs')):
            articles.append(normalize(candidate, str(path)))


def order(article):
    key = (article['slug'] + ' ' + article['title']).lower()
    rules = [('taylor', 0), ('jay', 1), ('leno', 1), ('vinny', 2), ('paz', 2), ('meta', 3), ('openai', 4), ('scialo', 5), ('venda', 6), ('ravioli', 6), ('matunuck', 7), ('raso', 7), ('iggy', 8), ('florence', 9)]
    for token, number in rules:
        if token in key:
            return number
    return 10


articles.sort(key=order)
if len(articles) != 10:
    raise SystemExit(f'Expected 10 stories, found {len(articles)}. No deliverable created.')
if len({a['slug'] for a in articles}) != 10:
    raise SystemExit('Duplicate article slug. No deliverable created.')

for a in articles:
    known = {str(s['id']) for s in a['source_notes']}
    for p in a['body']:
        if not isinstance(p['text'], str) or not p['text'].strip():
            raise ValueError(f'Empty paragraph in {a["slug"]}')
        if any(str(s) not in known for s in p['sources']):
            raise ValueError(f'Unresolved source in {a["slug"]}: {p["sources"]}')

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.72)
section.bottom_margin = Inches(.72)
section.left_margin = Inches(.82)
section.right_margin = Inches(.82)
section.footer_distance = Inches(.30)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Georgia'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Georgia')
    style.element.rPr.rFonts.set(qn('w:hAnsi'), 'Georgia')
    if style.element.pPr is not None:
        for child in list(style.element.pPr):
            if child.tag == qn('w:pBdr'):
                style.element.pPr.remove(child)

normal = doc.styles['Normal']
normal.font.size = Pt(11.5)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.widow_control = True
title = doc.styles['Title']
title.font.size = Pt(27)
title.font.bold = True
title.paragraph_format.space_after = Pt(12)
title.paragraph_format.keep_with_next = True
sub = doc.styles['Subtitle']
sub.font.size = Pt(12)
sub.font.italic = True
sub.paragraph_format.space_after = Pt(15)
sub.paragraph_format.keep_with_next = True
for name, size in [('Heading 1', 23), ('Heading 2', 16), ('Heading 3', 12)]:
    style = doc.styles[name]
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(10)
    style.paragraph_format.keep_with_next = True


def bookmark(p, name, number):
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(number))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(number))
    p._p.insert(0, start)
    p._p.append(end)


def link(p, label, target, external=True, size=10.5, color='000000'):
    h = OxmlElement('w:hyperlink')
    if external:
        rid = p.part.relate_to(target, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
        h.set(qn('r:id'), rid)
    else:
        h.set(qn('w:anchor'), target)
    r = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), color); props.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); props.append(u)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); props.append(sz)
    r.append(props)
    t = OxmlElement('w:t'); t.text = label; r.append(t)
    h.append(r); p._p.append(h)


def kicker(label):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(label.upper())
    r.font.name = 'Arial'
    r.font.size = Pt(9)
    r.bold = True
    return p


doc.add_paragraph('Rhode Island Hurricane', 'Title')
doc.add_paragraph('November 2026 stories and cover copy', 'Subtitle')
doc.add_paragraph('Ten article drafts for the first issue, covering celebrity stories, national technology news, Rhode Island food and a musical journey to Florence. Numbered references connect the reporting to the source list at the end.')
p = doc.add_paragraph('Prepared October 23 2026')
p.runs[0].font.size = Pt(10)

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [4.45, 1.6, .81]
for i, w in enumerate(widths):
    table.columns[i].width = Inches(w)
headers = ['Article', 'Section', 'Words']
for i, name in enumerate(headers):
    table.rows[0].cells[i].text = name
repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
for n, article in enumerate(articles, 1):
    cells = table.add_row().cells
    p = cells[0].paragraphs[0]
    link(p, article['title'], f'article_{n}', external=False, size=10)
    cells[1].text = article['section']
    cells[2].text = str(article['word_count'])

for row_number, row in enumerate(table.rows):
    for col, cell in enumerate(row.cells):
        cell.width = Inches(widths[col])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for side in ['top', 'left', 'bottom', 'right']:
            b = OxmlElement('w:' + side)
            b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9')
            borders.append(b)
        tcpr.append(borders)
        margins = OxmlElement('w:tcMar')
        for side in ['top', 'bottom', 'left', 'right']:
            m = OxmlElement('w:' + side); m.set(qn('w:w'), '95'); m.set(qn('w:type'), 'dxa'); margins.append(m)
        tcpr.append(margins)
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), '122948' if row_number == 0 else ('F3F5F7' if row_number % 2 == 0 else 'FFFFFF'))
        tcpr.append(shade)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.02
            if col == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(9.5)
                if row_number == 0:
                    run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)

cover_heading = doc.add_paragraph('Cover copy', 'Heading 1')
cover_heading.paragraph_format.page_break_before = True
doc.add_paragraph('November 2026', 'Subtitle')
cover_groups = [
    ('Masthead', 'RHODE ISLAND\nHURRICANE\nPEOPLE • POWER • CULTURE • TECH'),
    ('Main story', 'TAYLOR SWIFT’S $50,000 GIFT TO A HERO MOM\nShe stopped to help on a Rhode Island highway. It nearly cost her life.'),
    ('Supporting stories', 'JAY LENO MEETS THE BLACKBIRD\nA $2.5 million hypercar preview.\n\nVINNY PAZ THE COMEBACK THAT STILL HITS\nTen years after Bleed for This.\n\nRHODE ISLAND vs. META\nThe state takes on scam ads.\n\nOPENAI’S AGENTS BROKE CONTAINMENT\nInside the Hugging Face breach.'),
    ('Food line', 'RAVIOLI, CANNOLI & CLAM CAKES'),
    ('Model credit', 'COVER MODEL — IRINA VLADI\nIllustrative portrait. Irina is not a subject of the cover stories.')
]
for label, text in cover_groups:
    doc.add_paragraph(label, 'Heading 3')
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.keep_together = True

for n, article in enumerate(articles, 1):
    section_label = kicker(article['section'])
    section_label.paragraph_format.page_break_before = True
    p = doc.add_paragraph(article['title'], 'Heading 1')
    if article['slug'] == 'vinny-paz-ten-years':
        section_label.paragraph_format.space_after = Pt(4)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(8)
    bookmark(p, f'article_{n}', n)
    if article['dek']:
        dek_paragraph = doc.add_paragraph(article['dek'], 'Subtitle')
        if article['slug'] == 'vinny-paz-ten-years':
            dek_paragraph.paragraph_format.space_after = Pt(10)
    for item in article['body']:
        p = doc.add_paragraph(item['text'])
        if article['slug'] in {'jay-leno-blackbird', 'vinny-paz-ten-years'}:
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs:
                run.font.size = Pt(11)
            if article['slug'] == 'vinny-paz-ten-years':
                p.paragraph_format.space_after = Pt(2)
        if item['sources']:
            r = p.add_run(' ' + ','.join(str(x) for x in item['sources']))
            r.font.size = Pt(7.5)
            r.font.superscript = True

sources_heading = doc.add_paragraph('Sources', 'Heading 1')
sources_heading.paragraph_format.page_break_before = True
doc.add_paragraph('Source numbers restart for each article. Publication dates are given when available. Online listings and programs were checked on October 23 2026.')
source_lines = ['RHODE ISLAND HURRICANE', 'NOVEMBER 2026 SOURCE REGISTER', 'Prepared October 23, 2026', '']
for article in articles:
    doc.add_paragraph(article['title'], 'Heading 2')
    source_lines += [article['title'].upper(), '']
    for source_index, source in enumerate(article['source_notes']):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"{source['id']}. {source['outlet']} | {source['date']}")
        r.font.size = Pt(10)
        r.bold = True
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(11)
        p.paragraph_format.keep_with_next = source_index < len(article['source_notes']) - 1
        link(p, source['title'], source['url'])
        source_lines += [f"{source['id']}. {source['outlet']} | {source['date']}", source['title'], source['url'], '']

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('Rhode Island Hurricane | November 2026 | ')
r.font.name = 'Arial'; r.font.size = Pt(8)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
field_run = OxmlElement('w:r')
field_props = OxmlElement('w:rPr')
field_fonts = OxmlElement('w:rFonts')
field_fonts.set(qn('w:ascii'), 'Arial')
field_fonts.set(qn('w:hAnsi'), 'Arial')
field_props.append(field_fonts)
field_size = OxmlElement('w:sz')
field_size.set(qn('w:val'), '16')
field_props.append(field_size)
field_run.append(field_props)
field_text = OxmlElement('w:t')
field_text.text = '1'
field_run.append(field_text)
field.append(field_run)
footer._p.append(field)
doc.core_properties.title = 'Rhode Island Hurricane November 2026 Stories and Cover Copy'
doc.core_properties.subject = 'Ten article drafts and cover wording'
doc.core_properties.author = ''
doc.core_properties.comments = ''
doc.core_properties.keywords = 'Rhode Island Hurricane November 2026'
docx_path = OUT / 'Rhode-Island-Hurricane-November-2026-Stories.docx'
doc.save(docx_path)

for n, article in enumerate(articles, 1):
    lines = [article['title'].upper(), '', article['dek'], '']
    lines += [p['text'] + '\n' for p in article['body']]
    lines += ['SOURCES', '']
    for source in article['source_notes']:
        lines += [f"{source['id']}. {source['outlet']} | {source['date']}", source['title'], source['url'], '']
    (STORIES / f'{n:02d}-{article["slug"]}.txt').write_text('\n'.join(lines), encoding='utf-8')

(PACKAGE / 'Sources.txt').write_text('\n'.join(source_lines), encoding='utf-8')
shutil.copy2(OUT / 'Rhode-Island-Hurricane-Cover-Copy.txt', PACKAGE / 'Cover-Copy.txt')
direction = json.loads((ROOT / 'cover-direction.json').read_text())
notes = ['RHODE ISLAND HURRICANE', 'NOVEMBER 2026 EDITORIAL HANDOFF', '', 'COMPLETED', '',
         f"Ten story drafts totaling {sum(a['word_count'] for a in articles):,} body words, final proposed cover wording, a source register and this production handoff.", '',
         'COVER IMAGE STATUS', '', direction['status'], '', 'ART DIRECTION', '', direction['reference'], '']
for element in direction['hierarchy']:
    notes += [element['element'].upper() + ' - ' + element['placement'], element['instruction'], '']
notes += ['COVER PRODUCTION', ''] + [p + '\n' for p in direction['production_notes']]
notes += ['ARTICLE HANDOFF NOTES', '',
          'These pieces are based on identified published interviews, records and official materials. No original interview, restaurant visit or tasting has been represented as having occurred. Photo sourcing and a final printed page design are separate production work.', '']
for article in articles:
    notes += [article['title'].upper()]
    limits = article.get('reporting_limits', [])
    if isinstance(limits, str):
        limits = [limits]
    for item in limits:
        notes.append('- ' + (item if isinstance(item, str) else json.dumps(item, ensure_ascii=False)))
    notes += ['']
notes += ['PUBLICATION STATUS', '', 'No website was changed and nothing was published or sent to a third party.']
(PACKAGE / 'Editorial-Handoff.txt').write_text('\n'.join(notes), encoding='utf-8')
(PACKAGE / 'Start-Here.txt').write_text(
    'RHODE ISLAND HURRICANE\nNOVEMBER 2026\n\n'
    'Open the Word manuscript to read the complete issue text and proposed cover copy.\n'
    'The Stories folder contains one copyable text file for each article, with its sources at the bottom.\n'
    'Cover-Copy.txt contains only the proposed cover wording.\n'
    'Editorial-Handoff.txt records art direction, article update notes and the unfinished cover-image status.\n'
    'Sources.txt contains every source URL in plain text.\n', encoding='utf-8')
shutil.copy2(docx_path, PACKAGE / docx_path.name)
zip_path = OUT / 'Rhode-Island-Hurricane-November-2026-Editorial-Package.zip'
with zipfile.ZipFile(zip_path, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
    for path in sorted(PACKAGE.rglob('*')):
        if path.is_file():
            zf.write(path, path.relative_to(OUT))
manifest = {
    'article_count': len(articles),
    'body_word_count': sum(a['word_count'] for a in articles),
    'articles': [{'title': a['title'], 'slug': a['slug'], 'section': a['section'], 'words': a['word_count']} for a in articles],
    'outputs': [str(docx_path), str(OUT / 'Rhode-Island-Hurricane-Cover-Copy.txt'), str(zip_path)]
}
(ROOT / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
(ROOT / 'normalized-articles.json').write_text(json.dumps(articles, ensure_ascii=False, indent=2))
print(json.dumps(manifest, ensure_ascii=False, indent=2))
